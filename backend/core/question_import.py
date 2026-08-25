"""
题目批量导入的共享逻辑，支持 .xlsx 和 .csv。
字段顺序: 序号,题型,题目,选项A,选项B,选项C,选项D,答案,标签

被 core/management/commands/import_questions_excel.py（命令行）
和 core/admin.py（Django admin 页面上传）共用。
"""
import csv
import io

import openpyxl
from django.db import transaction
from core.models import Tag, Question

QTYPE_MAP = {
    '单选': Question.QuestionType.SINGLE,
    '多选': Question.QuestionType.MULTIPLE,
    '判断题': Question.QuestionType.TRUE_FALSE,
}

OPTION_KEYS = ['A', 'B', 'C', 'D']

# 已知的原始数据笔误修正：第116题答案 'AR' 应为多选 'AB'（选项无 R，用户已确认）
ANSWER_OVERRIDES = {
    116: 'AB',
}


class ImportResult:
    def __init__(self):
        self.created_count = 0
        self.skipped_dup_in_file = 0
        self.skipped_existing_in_db = 0
        self.skipped_invalid = 0
        self.missing_tags = set()
        self.messages = []  # (level, text) 列表，level: 'warning' | 'error'

    def warn(self, text):
        self.messages.append(('warning', text))

    def error(self, text):
        self.messages.append(('error', text))

    @property
    def summary(self):
        return (
            f'导入 {self.created_count} 条，'
            f'文件内重复跳过 {self.skipped_dup_in_file} 条，'
            f'数据库已存在跳过 {self.skipped_existing_in_db} 条，'
            f'无效数据跳过 {self.skipped_invalid} 条'
        )


def import_questions_from_workbook(file_obj, dry_run=False):
    """
    file_obj: 文件路径字符串，或任何 openpyxl.load_workbook 接受的文件对象/句柄。
    dry_run: True 时只校验，不写入数据库（事务会回滚）。
    返回 ImportResult。
    """
    wb = openpyxl.load_workbook(file_obj, data_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))[1:]
    return _import_rows(rows, dry_run=dry_run)


def import_questions_from_csv(file_obj, dry_run=False):
    """
    file_obj: 以文本模式打开的文件对象，或字节内容（会按 utf-8-sig 解码）。
    dry_run: True 时只校验，不写入数据库（事务会回滚）。
    返回 ImportResult。
    """
    if isinstance(file_obj, bytes):
        text = file_obj.decode('utf-8-sig')
    elif hasattr(file_obj, 'read'):
        content = file_obj.read()
        text = content.decode('utf-8-sig') if isinstance(content, bytes) else content
    else:
        text = file_obj

    reader = csv.reader(io.StringIO(text))
    all_rows = list(reader)[1:]  # 跳过表头

    rows = []
    for r in all_rows:
        r = r + [None] * (9 - len(r))  # 补齐到9列，缺列的单元格视为空
        seq_raw = r[0]
        try:
            seq = int(seq_raw)
        except (TypeError, ValueError):
            seq = seq_raw
        rows.append((seq, r[1], r[2], r[3], r[4], r[5], r[6], r[7], r[8]))

    return _import_rows(rows, dry_run=dry_run)


def _import_rows(rows, dry_run=False):
    result = ImportResult()
    seen_contents = set()

    with transaction.atomic():
        for row in rows:
            _process_row(row, result, seen_contents, dry_run)

        if dry_run:
            transaction.set_rollback(True)

    return result


def _process_row(row, result, seen_contents, dry_run):
    seq, qtype_cn, content, opt_a, opt_b, opt_c, opt_d, answer, tags_cn = row[:9]

    if not content or not str(content).strip():
        result.skipped_invalid += 1
        result.warn(f'第{seq}行: 题目内容为空，跳过')
        return
    content = str(content).strip()

    if content in seen_contents:
        result.skipped_dup_in_file += 1
        result.warn(f'第{seq}行: 文件内重复题目，跳过 -> {content[:30]}...')
        return
    seen_contents.add(content)

    if Question.objects.filter(content=content).exists():
        result.skipped_existing_in_db += 1
        result.warn(f'第{seq}行: 数据库中已存在同内容题目，跳过 -> {content[:30]}...')
        return

    qtype = QTYPE_MAP.get(str(qtype_cn).strip())
    if qtype is None:
        result.skipped_invalid += 1
        result.error(f'第{seq}行: 未知题型 "{qtype_cn}"，跳过')
        return

    raw_answer = ANSWER_OVERRIDES.get(seq, answer)
    raw_answer = str(raw_answer).strip() if raw_answer is not None else ''

    if qtype == Question.QuestionType.TRUE_FALSE:
        options = [{"key": "True", "text": "正确"}, {"key": "False", "text": "错误"}]
        if raw_answer == '正确':
            correct_answer = 'True'
        elif raw_answer == '错误':
            correct_answer = 'False'
        else:
            result.skipped_invalid += 1
            result.error(f'第{seq}行: 判断题答案无法识别 "{raw_answer}"，跳过')
            return
    else:
        options = []
        for key, text in zip(OPTION_KEYS, [opt_a, opt_b, opt_c, opt_d]):
            if text is not None and str(text).strip():
                options.append({"key": key, "text": str(text).strip()})

        valid_keys = {o['key'] for o in options}
        answer_letters = [c for c in raw_answer if c.strip()]
        if not answer_letters or any(c not in valid_keys for c in answer_letters):
            result.skipped_invalid += 1
            result.error(f'第{seq}行: 答案 "{raw_answer}" 超出选项范围 {sorted(valid_keys)}，跳过')
            return

        if qtype == Question.QuestionType.SINGLE:
            correct_answer = answer_letters[0]
        else:
            correct_answer = ','.join(answer_letters)

    tag_names = [t.strip() for t in str(tags_cn or '').split('、') if t.strip()]
    tag_objs = []
    for name in tag_names:
        tag = Tag.objects.filter(name=name).first()
        if tag is None:
            result.missing_tags.add(name)
        else:
            tag_objs.append(tag)

    if dry_run:
        result.created_count += 1
        return

    question = Question.objects.create(
        content=content,
        question_type=qtype,
        options=options,
        correct_answer=correct_answer,
        difficulty=3,
        is_active=True,
    )
    question.tags.set(tag_objs)
    result.created_count += 1
